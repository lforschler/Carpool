#!/usr/bin/env python3
"""Import the Google Sheet '2027' tab into src/lib/mock-schedule.json.

Source sheet:
  https://docs.google.com/spreadsheets/d/1A_w9-kkoFYO_WLQZaS2KQ6lsJ6pn3nsExurtwoBtems/edit
Tab name: 2027 (gid=0). Fall semester; weekdays align to calendar year 2026
(school year 2026–27). The '2026' tab is spring and is intentionally ignored.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import tempfile
import urllib.request
from datetime import date, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl required: pip install openpyxl\n")
    raise

SHEET_ID = "1A_w9-kkoFYO_WLQZaS2KQ6lsJ6pn3nsExurtwoBtems"
SHEET_TAB = "2027"
XLSX_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
SOURCE_URL = (
    f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"
    f"?usp=drivesdk#gid=0"
)

WEEKDAY_NAMES = [
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
    "Sunday",
]

DRIVER_MAP = {
    "lucas": "drv-lucas",
    "dee": "drv-dee",
    "jill": "drv-jill",
    "nicola": "drv-nicola",
}

CORE_KIDS = ["kid-eli", "kid-emily", "kid-blake", "kid-anderson"]
KID_ALIASES = {
    "a": "kid-anderson",
    "anderson": "kid-anderson",
    "eli": "kid-eli",
    "emily": "kid-emily",
    "blake": "kid-blake",
    "nico": "kid-nico",
    "andy": "kid-andy",
}

FAMILIES = [
    {
        "id": "fam-jill",
        "name": "Weis",
        "parentName": "Jill Weis",
        "phone": "417-612-3768",
        "address": "5355 S. Westwood Ave",
        "driverId": "drv-jill",
        "kidIds": ["kid-eli", "kid-emily"],
    },
    {
        "id": "fam-lucas",
        "name": "Forschler",
        "parentName": "Lucas Forschler",
        "phone": "408-368-5949",
        "address": "4830 S. Wellington Drive",
        "driverId": "drv-lucas",
        "kidIds": ["kid-blake"],
    },
    {
        "id": "fam-dee",
        "name": "Brewer",
        "parentName": "Dee Brewer",
        "phone": "618-367-1329",
        "address": "3132 W. Wick St",
        "driverId": "drv-dee",
        "kidIds": ["kid-anderson"],
    },
]

KIDS = [
    {"id": "kid-eli", "name": "Eli", "familyId": "fam-jill", "short": "Eli"},
    {"id": "kid-emily", "name": "Emily", "familyId": "fam-jill", "short": "Emily"},
    {"id": "kid-blake", "name": "Blake", "familyId": "fam-lucas", "short": "Blake"},
    {
        "id": "kid-anderson",
        "name": "Anderson",
        "familyId": "fam-dee",
        "short": "A",
    },
    {
        "id": "kid-nico",
        "name": "Nico",
        "familyId": None,
        "short": "Nico",
        "note": "Guest rider from sheet",
    },
    {
        "id": "kid-andy",
        "name": "Andy",
        "familyId": None,
        "short": "Andy",
        "note": "Guest rider from sheet",
    },
]

DRIVERS = [
    {
        "id": "drv-jill",
        "name": "Jill",
        "fullName": "Jill Weis",
        "familyId": "fam-jill",
        "seatCapacity": 5,
    },
    {
        "id": "drv-lucas",
        "name": "Lucas",
        "fullName": "Lucas Forschler",
        "familyId": "fam-lucas",
        "seatCapacity": 5,
    },
    {
        "id": "drv-dee",
        "name": "Dee",
        "fullName": "Dee Brewer",
        "familyId": "fam-dee",
        "seatCapacity": 5,
    },
    {
        "id": "drv-nicola",
        "name": "Nicola",
        "fullName": "Nicola",
        "familyId": None,
        "seatCapacity": 5,
    },
]


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def normalize_note(raw: str) -> str:
    s = raw.replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r" *\n *", "\n", s).strip()
    # Fix common sheet typos
    s = s.replace("Lucas No A)", "Lucas (No A)")
    s = re.sub(r"\s+", " ", s)
    return s


def parse_cell(raw: str) -> dict:
    note = normalize_note(raw)
    if not note:
        return {"empty": True}

    lower = note.lower()
    if "no school" in lower:
        return {
            "empty": False,
            "noSchool": True,
            "driverId": None,
            "riderIds": [],
            "excludedIds": [],
            "status": "cancelled",
            "sheetNote": "No School",
        }

    driver_id = None
    for name, did in DRIVER_MAP.items():
        if re.search(rf"\b{name}\b", lower):
            driver_id = did
            break

    excluded: list[str] = []
    # (No A), (no A), (No A or Blake), (no eli & Emily), etc.
    for m in re.finditer(r"\(\s*no\s+([^)]+)\)", lower):
        chunk = m.group(1)
        chunk = chunk.replace(" or ", " ").replace("&", " ").replace(",", " ")
        for token in chunk.split():
            token = token.strip()
            if token in KID_ALIASES:
                kid = KID_ALIASES[token]
                if kid not in excluded:
                    excluded.append(kid)

    guests: list[str] = []
    for guest_name, kid_id in (("nico", "kid-nico"), ("andy", "kid-andy")):
        if re.search(rf"\b{guest_name}\b", lower):
            guests.append(kid_id)

    riders = [k for k in CORE_KIDS if k not in excluded] + [
        g for g in guests if g not in excluded
    ]

    status = "unassigned" if driver_id is None else "planned"
    return {
        "empty": False,
        "noSchool": False,
        "driverId": driver_id,
        "riderIds": riders,
        "excludedIds": excluded,
        "status": status,
        "sheetNote": note,
    }


def extract_week_blocks(ws) -> list[dict]:
    """Extract per-column-group day cells with AM/PM notes.

    The sheet lays out four month streams side-by-side. Each stream must be
    dated independently (left-to-right across a row is NOT chronological).
    """
    max_row = min(ws.max_row, 80)
    max_col = min(ws.max_column, 28)

    grid = [
        [ws.cell(r, c).value for c in range(1, max_col + 1)]
        for r in range(1, max_row + 1)
    ]

    # 0-based day columns for each month stream (Mon–Fri)
    streams = [
        {"id": 0, "day_cols": [1, 2, 3, 4, 5], "label_cols": [0]},
        {"id": 1, "day_cols": [8, 9, 10, 11, 12], "label_cols": [7]},
        {"id": 2, "day_cols": [15, 16, 17, 18, 19], "label_cols": [14]},
        {"id": 3, "day_cols": [22, 23, 24, 25, 26], "label_cols": [21]},
    ]

    blocks: list[dict] = []
    r = 0
    while r < len(grid) - 3:
        row = grid[r]
        day_hits = 0
        for stream in streams:
            for c in stream["day_cols"]:
                if c < len(row) and isinstance(row[c], (int, float)) and 1 <= float(row[c]) <= 31:
                    day_hits += 1
        if day_hits < 3:
            r += 1
            continue

        weekday_row = grid[r + 1] if r + 1 < len(grid) else []
        am_row = pm_row = None
        for k in range(2, 6):
            if r + k >= len(grid):
                break
            labels = {
                cell_str(grid[r + k][c]).upper()
                for c in (0, 7, 14, 21)
                if c < len(grid[r + k])
            }
            if am_row is None and any(x.startswith("AM") for x in labels):
                am_row = grid[r + k]
            if pm_row is None and any(x == "PM" or x.startswith("PM") for x in labels):
                pm_row = grid[r + k]
        if am_row is None or pm_row is None:
            r += 1
            continue

        for stream in streams:
            for c in stream["day_cols"]:
                day_val = row[c] if c < len(row) else None
                wd = cell_str(weekday_row[c]) if c < len(weekday_row) else ""
                if not isinstance(day_val, (int, float)):
                    continue
                day = int(day_val)
                if wd not in WEEKDAY_NAMES[:5]:
                    continue
                am_raw = cell_str(am_row[c]) if c < len(am_row) else ""
                pm_raw = cell_str(pm_row[c]) if c < len(pm_row) else ""
                if am_raw.upper().startswith("AM") and len(am_raw.strip()) <= 3:
                    am_raw = ""
                if pm_raw.upper() == "PM":
                    pm_raw = ""
                blocks.append(
                    {
                        "stream": stream["id"],
                        "day": day,
                        "weekday": wd,
                        "am": am_raw,
                        "pm": pm_raw,
                        "row": r + 1,
                        "col": c + 1,
                    }
                )
        r += 2

    return blocks


def next_matching_date(day: int, weekday_name: str, on_or_after: date) -> date | None:
    """Next calendar date with this day-of-month and weekday, on/after on_or_after."""
    want_wd = WEEKDAY_NAMES.index(weekday_name)
    for i in range(0, 370):
        d = on_or_after + timedelta(days=i)
        if d.day == day and d.weekday() == want_wd:
            return d
    return None


def build_trips(blocks: list[dict], start_year: int = 2026) -> list[dict]:
    # Each side-by-side month column is an independent chronological stream.
    stream_start = {
        0: date(start_year, 8, 1),
        1: date(start_year, 9, 1),
        2: date(start_year, 10, 1),
        3: date(start_year, 11, 1),
    }

    trips: list[dict] = []
    seen: set[tuple[str, str]] = set()

    for stream_id in sorted(stream_start):
        cursor = stream_start[stream_id]
        stream_blocks = [b for b in blocks if b["stream"] == stream_id]
        stream_blocks.sort(key=lambda b: (b["row"], b["col"]))
        for b in stream_blocks:
            d = next_matching_date(b["day"], b["weekday"], cursor)
            if d is None:
                print(f"WARN: could not resolve {b}", file=sys.stderr)
                continue
            cursor = d + timedelta(days=1)

            if d < date(start_year, 8, 1) or d > date(start_year, 12, 31):
                continue

            for direction, raw, time in (
                ("AM", b["am"], "7:30 AM"),
                ("PM", b["pm"], "3:15 PM"),
            ):
                parsed = parse_cell(raw)
                if parsed.get("empty"):
                    continue
                key = (d.isoformat(), direction)
                if key in seen:
                    print(
                        f"WARN: duplicate {key} from r{b['row']}c{b['col']} "
                        f"(keeping first)",
                        file=sys.stderr,
                    )
                    continue
                seen.add(key)
                trips.append(
                    {
                        "id": f"trip-{d.isoformat()}-{direction.lower()}",
                        "date": d.isoformat(),
                        "direction": direction,
                        "time": time,
                        "driverId": parsed["driverId"],
                        "riderIds": parsed["riderIds"],
                        "excludedIds": parsed["excludedIds"],
                        "status": parsed["status"],
                        "noSchool": parsed["noSchool"],
                        "sheetNote": parsed["sheetNote"],
                    }
                )

    trips.sort(key=lambda t: (t["date"], 0 if t["direction"] == "AM" else 1))
    return trips


def download_xlsx(dest: Path) -> Path:
    print(f"Downloading {XLSX_URL}")
    urllib.request.urlretrieve(XLSX_URL, dest)
    return dest


def main() -> int:
    repo = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--xlsx",
        type=Path,
        help="Local xlsx path (default: download from Google Sheets)",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=repo / "src/lib/mock-schedule.json",
    )
    parser.add_argument("--start-year", type=int, default=2026)
    args = parser.parse_args()

    if args.xlsx:
        xlsx_path = args.xlsx
    else:
        tmp = Path(tempfile.mkdtemp()) / "carpool-schedule.xlsx"
        xlsx_path = download_xlsx(tmp)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET_TAB not in wb.sheetnames:
        print(f"Missing tab {SHEET_TAB!r}; found {wb.sheetnames}", file=sys.stderr)
        return 1
    ws = wb[SHEET_TAB]
    blocks = extract_week_blocks(ws)
    trips = build_trips(blocks, start_year=args.start_year)

    if not trips:
        print("No trips parsed", file=sys.stderr)
        return 1

    dates = [t["date"] for t in trips]
    out = {
        "source": SOURCE_URL,
        "sheetTab": SHEET_TAB,
        "schoolYearLabel": (
            f"Fall {args.start_year} / school year {args.start_year}–"
            f"{args.start_year + 1} (Google Sheet tab '{SHEET_TAB}')"
        ),
        "families": FAMILIES,
        "kids": KIDS,
        "drivers": DRIVERS,
        "trips": trips,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(out, indent=2) + "\n")
    print(f"Wrote {args.out}")
    print(f"Trips: {len(trips)}")
    print(f"Date range: {min(dates)} → {max(dates)}")
    print(f"No-school trips: {sum(1 for t in trips if t['noSchool'])}")
    print(f"Unassigned: {sum(1 for t in trips if t['status'] == 'unassigned')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
